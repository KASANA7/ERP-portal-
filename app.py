from flask import Flask, render_template, request, redirect, url_for, session, jsonify, send_file, flash
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from database import get_db, init_db, generate_card_number
from datetime import datetime, timedelta
import os, json, io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = 'cms-moon-beverages-dasna-2026'

# Custom Jinja filter
import json as _json
app.jinja_env.filters['fromjson'] = lambda s: _json.loads(s) if s else []

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max

ALLOWED_EXTENSIONS = {'pdf', 'jpg', 'jpeg', 'png'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def save_file(file, subfolder='docs'):
    if file and allowed_file(file.filename):
        folder = os.path.join(app.config['UPLOAD_FOLDER'], subfolder)
        os.makedirs(folder, exist_ok=True)
        filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{secure_filename(file.filename)}"
        path = os.path.join(folder, filename)
        file.save(path)
        return os.path.join(subfolder, filename)
    return None

# ── AUTH ─────────────────────────────────────────────────
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        db = get_db()
        user = db.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
        db.close()
        if user and check_password_hash(user['password_hash'], password):
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['full_name'] = user['full_name']
            session['role'] = user['role']
            return redirect(url_for('dashboard'))
        return render_template('login.html', error='Invalid username or password')
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ── DASHBOARD ────────────────────────────────────────────
@app.route('/')
@login_required
def dashboard():
    db = get_db()
    stats = {
        'companies': db.execute("SELECT COUNT(*) FROM companies").fetchone()[0],
        'companies_pending': db.execute("SELECT COUNT(*) FROM companies WHERE status='Pending'").fetchone()[0],
        'supervisors': db.execute("SELECT COUNT(*) FROM supervisors").fetchone()[0],
        'cards_total': db.execute("SELECT COUNT(*) FROM worker_cards").fetchone()[0],
        'cards_active': db.execute("SELECT COUNT(*) FROM worker_cards WHERE status='Approved'").fetchone()[0],
        'cards_pending': db.execute("SELECT COUNT(*) FROM worker_cards WHERE status='Pending'").fetchone()[0],
        'cards_today': db.execute("SELECT COUNT(*) FROM worker_cards WHERE DATE(created_at)=DATE('now')").fetchone()[0],
    }
    # Last 7 days chart data
    chart = []
    for i in range(6, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime('%Y-%m-%d')
        label = (datetime.now() - timedelta(days=i)).strftime('%a')
        count = db.execute("SELECT COUNT(*) FROM worker_cards WHERE DATE(created_at)=?", (d,)).fetchone()[0]
        chart.append({'day': label, 'count': count})
    recent_cards = db.execute('''
        SELECT wc.*, c.name as company_name
        FROM worker_cards wc
        LEFT JOIN companies c ON wc.company_id = c.id
        ORDER BY wc.created_at DESC LIMIT 5
    ''').fetchall()
    pending_approvals = db.execute('''
        SELECT wc.*, c.name as company_name
        FROM worker_cards wc
        LEFT JOIN companies c ON wc.company_id = c.id
        WHERE wc.status='Pending'
        ORDER BY wc.created_at DESC LIMIT 5
    ''').fetchall()
    db.close()
    return render_template('dashboard.html', stats=stats, chart=chart,
                           recent_cards=recent_cards, pending_approvals=pending_approvals)

# ── COMPANIES ────────────────────────────────────────────
@app.route('/companies')
@login_required
def companies():
    db = get_db()
    companies = db.execute('SELECT * FROM companies ORDER BY created_at DESC').fetchall()
    db.close()
    return render_template('companies.html', companies=companies)

@app.route('/companies/add', methods=['GET', 'POST'])
@login_required
def add_company():
    if request.method == 'POST':
        # Backend guard — reject if form wasn't completed properly
        if request.form.get('step_completed') != '5':
            flash('Please complete all steps before submitting.', 'error')
            return redirect(url_for('add_company'))

        # Also validate required fields server-side
        required = ['name', 'contact_person', 'contact_number', 'address',
                    'work_type', 'num_supervisors', 'num_workers', 'contract_start',
                    'nature_of_work']
        for field in required:
            if not request.form.get(field, '').strip():
                flash(f'Required field missing: {field.replace("_"," ").title()}', 'error')
                return redirect(url_for('add_company'))
        db = get_db()

        # Part A checklist (HR - 5 questions)
        parta = {}
        for key in request.form:
            if key.startswith('parta_'):
                parta[key] = {
                    'status': request.form.get(key),
                    'remark': request.form.get(key + '_remark', '')
                }

        # Part B checklist (Safety - 20 questions)
        partb = {}
        for key in request.form:
            if key.startswith('partb_'):
                partb[key] = {
                    'status': request.form.get(key),
                    'remark': request.form.get(key + '_remark', '')
                }

        # Calculate expiry (1 year from start)
        contract_start = request.form.get('contract_start')
        contract_expiry = None
        if contract_start:
            try:
                start_dt = datetime.strptime(contract_start, '%Y-%m-%d')
                contract_expiry = (start_dt + timedelta(days=365)).strftime('%Y-%m-%d')
            except:
                pass

        num_workers = int(request.form.get('num_workers', 0) or 0)

        db.execute('''INSERT INTO companies
            (name, gst, contact_person, contact_number, email, esic_pf_code,
             address, work_type, num_supervisors, num_workers, risk_level,
             contract_start, contract_expiry, nature_of_work,
             parta_checklist, partb_checklist)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
            request.form.get('name'),
            request.form.get('gst'),
            request.form.get('contact_person'),
            request.form.get('contact_number'),
            request.form.get('email'),
            request.form.get('esic_pf_code'),
            request.form.get('address'),
            request.form.get('work_type'),
            request.form.get('num_supervisors', 0),
            num_workers,
            request.form.get('risk_level'),
            contract_start,
            contract_expiry,
            request.form.get('nature_of_work'),
            json.dumps(parta),
            json.dumps(partb)
        ))
        company_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]

        # Part A documents
        parta_names = request.form.getlist('parta_doc_name[]')
        parta_files = request.files.getlist('parta_doc_file[]')
        for name, file in zip(parta_names, parta_files):
            path = save_file(file, 'companies')
            if path:
                db.execute('INSERT INTO company_docs (company_id, doc_name, file_path, part) VALUES (?,?,?,?)',
                           (company_id, name, path, 'A'))

        # Part B documents
        partb_names = request.form.getlist('partb_doc_name[]')
        partb_files = request.files.getlist('partb_doc_file[]')
        for name, file in zip(partb_names, partb_files):
            path = save_file(file, 'companies')
            if path:
                db.execute('INSERT INTO company_docs (company_id, doc_name, file_path, part) VALUES (?,?,?,?)',
                           (company_id, name, path, 'B'))

        db.commit()
        db.close()
        flash('Company submitted for approval successfully!', 'success')
        return redirect(url_for('companies'))
    return render_template('add_company.html')

@app.route('/companies/<int:company_id>')
@login_required
def view_company(company_id):
    db = get_db()
    company = db.execute('SELECT * FROM companies WHERE id=?', (company_id,)).fetchone()
    if not company:
        db.close()
        flash('Company not found', 'error')
        return redirect(url_for('companies'))
    docs_a = db.execute("SELECT * FROM company_docs WHERE company_id=? AND part='A'", (company_id,)).fetchall()
    docs_b = db.execute("SELECT * FROM company_docs WHERE company_id=? AND part='B'", (company_id,)).fetchall()
    supervisors = db.execute("SELECT * FROM supervisors WHERE company_id=?", (company_id,)).fetchall()
    cards_count = db.execute("SELECT COUNT(*) FROM worker_cards WHERE company_id=?", (company_id,)).fetchone()[0]
    db.close()

    # Parse checklists
    parta = {}
    partb = {}
    try:
        parta = json.loads(company['parta_checklist'] or '{}')
    except: pass
    try:
        partb = json.loads(company['partb_checklist'] or '{}')
    except: pass

    # Expiry status
    expiry_status = 'ok'
    days_left = None
    if company['contract_expiry']:
        try:
            exp = datetime.strptime(company['contract_expiry'], '%Y-%m-%d')
            days_left = (exp - datetime.now()).days
            if days_left < 0:
                expiry_status = 'expired'
            elif days_left <= 30:
                expiry_status = 'warning'
        except: pass

    return render_template('view_company.html',
        company=company, docs_a=docs_a, docs_b=docs_b,
        supervisors=supervisors, cards_count=cards_count,
        parta=parta, partb=partb,
        expiry_status=expiry_status, days_left=days_left,
        now=datetime.now().strftime('%d.%b.%Y'))

@app.route('/companies/<int:company_id>/toggle-active', methods=['POST'])
@login_required
def toggle_active(company_id):
    db = get_db()
    co = db.execute("SELECT is_active FROM companies WHERE id=?", (company_id,)).fetchone()
    new_val = 0 if co['is_active'] else 1
    db.execute("UPDATE companies SET is_active=? WHERE id=?", (new_val, company_id))
    db.commit()
    db.close()
    return jsonify({'success': True, 'is_active': new_val})

@app.route('/companies/save-draft', methods=['POST'])
@login_required
def save_draft():
    data = request.json
    company_id = data.get('company_id')
    step = data.get('step', 1)
    form_data = data.get('form_data', {})

    # Build checklist JSON from form_data
    parta = {}
    partb = {}
    for key, val in form_data.items():
        if key.startswith('parta_') and not key.endswith('_remark'):
            parta[key] = {'status': val, 'remark': form_data.get(key+'_remark','')}
        if key.startswith('partb_') and not key.endswith('_remark'):
            partb[key] = {'status': val, 'remark': form_data.get(key+'_remark','')}

    db = get_db()
    common_fields = (
        form_data.get('name') or 'Draft',
        step,
        json.dumps(form_data),
        form_data.get('gst',''),
        form_data.get('contact_person',''),
        form_data.get('contact_number',''),
        form_data.get('email',''),
        form_data.get('esic_pf_code',''),
        form_data.get('address',''),
        form_data.get('work_type',''),
        int(form_data.get('num_supervisors') or 0),
        int(form_data.get('num_workers') or 0),
        form_data.get('contract_start',''),
        form_data.get('nature_of_work',''),
        form_data.get('risk_level',''),
    )

    if company_id:
        db.execute("""UPDATE companies SET
            name=?, draft_step=?, draft_data=?,
            gst=?, contact_person=?, contact_number=?,
            email=?, esic_pf_code=?, address=?,
            work_type=?, num_supervisors=?, num_workers=?,
            contract_start=?, nature_of_work=?, risk_level=?,
            parta_checklist=CASE WHEN ? != '{}' THEN ? ELSE parta_checklist END,
            partb_checklist=CASE WHEN ? != '{}' THEN ? ELSE partb_checklist END
            WHERE id=?""", (
            *common_fields,
            json.dumps(parta), json.dumps(parta),
            json.dumps(partb), json.dumps(partb),
            company_id
        ))
    else:
        db.execute("""INSERT INTO companies
            (name, draft_step, draft_data,
             gst, contact_person, contact_number,
             email, esic_pf_code, address,
             work_type, num_supervisors, num_workers,
             contract_start, nature_of_work, risk_level,
             parta_checklist, partb_checklist, status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Draft')""", (
            *common_fields,
            json.dumps(parta) or None,
            json.dumps(partb) or None,
        ))
        company_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]

    db.commit()
    db.close()
    return jsonify({'success': True, 'company_id': company_id})

@app.route('/companies/<int:company_id>/continue', methods=['GET'])
@login_required
def continue_draft(company_id):
    db = get_db()
    co = db.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
    db.close()
    if not co or co['status'] != 'Draft':
        return redirect(url_for('companies'))
    draft_data = {}
    try:
        draft_data = json.loads(co['draft_data'] or '{}')
    except: pass
    return render_template('add_company.html',
        draft=co, draft_data=draft_data,
        draft_step=co['draft_step'] or 1)

@app.route('/companies/<int:company_id>/activate', methods=['POST'])
@login_required
def activate_company(company_id):
    db = get_db()
    db.execute("UPDATE companies SET is_active=1 WHERE id=?", (company_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/companies/<int:company_id>/deactivate', methods=['POST'])
@login_required
def deactivate_company(company_id):
    db = get_db()
    db.execute("UPDATE companies SET is_active=0 WHERE id=?", (company_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/companies/<int:company_id>/approve', methods=['POST'])
@login_required
def approve_company(company_id):
    data = request.json or {}
    part = data.get('part', 'safety')
    remarks = data.get('remarks', 'Approved')
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    approver = session.get('full_name') or session.get('username')
    db = get_db()
    if part == 'hr':
        db.execute("""UPDATE companies SET
            hr_status='Approved', hr_approved_by=?, hr_approved_at=?, hr_remarks=?
            WHERE id=?""", (approver, now, remarks, company_id))
    elif part == 'cluster':
        db.execute("""UPDATE companies SET
            cluster_approved_by=?, cluster_approved_at=?, cluster_remarks=?
            WHERE id=?""", (approver, now, remarks, company_id))
    else:
        db.execute("""UPDATE companies SET
            safety_status='Approved', safety_approved_by=?, safety_approved_at=?, safety_remarks=?
            WHERE id=?""", (approver, now, remarks, company_id))
    # Check if both HR and Safety approved → set overall Approved
    co = db.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
    if co['hr_status'] == 'Approved' and co['safety_status'] == 'Approved':
        db.execute("UPDATE companies SET status='Approved' WHERE id=?", (company_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/companies/<int:company_id>/reject', methods=['POST'])
@login_required
def reject_company(company_id):
    data = request.json or {}
    part = data.get('part', 'safety')
    remarks = data.get('remarks', 'Rejected')
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    approver = session.get('full_name') or session.get('username')
    db = get_db()
    if part == 'hr':
        db.execute("""UPDATE companies SET
            hr_status='Rejected', hr_approved_by=?, hr_approved_at=?, hr_remarks=?
            WHERE id=?""", (approver, now, remarks, company_id))
    else:
        db.execute("""UPDATE companies SET
            safety_status='Rejected', safety_approved_by=?, safety_approved_at=?, safety_remarks=?
            WHERE id=?""", (approver, now, remarks, company_id))
    db.execute("UPDATE companies SET status='Rejected' WHERE id=?", (company_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/supervisors/save-draft', methods=['POST'])
@login_required
def save_sup_draft():
    data = request.json
    supervisor_id = data.get('supervisor_id')
    form_data = data.get('form_data', {})
    db = get_db()
    fields = (
        form_data.get('first_name', 'Draft'),
        form_data.get('middle_name', ''),
        form_data.get('last_name', ''),
        form_data.get('dob', ''),
        form_data.get('gender', 'Male'),
        form_data.get('mobile', ''),
        form_data.get('aadhar', ''),
        form_data.get('address', ''),
        form_data.get('languages', ''),
        form_data.get('qualifications', ''),
        form_data.get('work_type', ''),
        int(form_data.get('num_workers') or 0),
        int(form_data.get('experience_years') or 0),
        int(form_data.get('experience_months') or 0),
        form_data.get('technical_qualification', ''),
        form_data.get('jobs_handled', ''),
        form_data.get('risk_level', ''),
        form_data.get('nature_of_work', ''),
        form_data.get('remarks', ''),
        json.dumps(form_data),
        form_data.get('company_id') or None,
    )
    if supervisor_id:
        db.execute("""UPDATE supervisors SET
            first_name=?, middle_name=?, last_name=?, dob=?, gender=?,
            mobile=?, aadhar=?, address=?, languages=?, qualifications=?,
            work_type=?, num_workers=?, experience_years=?, experience_months=?,
            technical_qualification=?, jobs_handled=?, risk_level=?,
            nature_of_work=?, remarks=?, draft_data=?, company_id=?
            WHERE id=?""", (*fields, supervisor_id))
    else:
        db.execute("""INSERT INTO supervisors
            (first_name, middle_name, last_name, dob, gender,
             mobile, aadhar, address, languages, qualifications,
             work_type, num_workers, experience_years, experience_months,
             technical_qualification, jobs_handled, risk_level,
             nature_of_work, remarks, draft_data, company_id, status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Draft')""", fields)
        supervisor_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]
    db.commit()
    db.close()
    return jsonify({'success': True, 'supervisor_id': supervisor_id})

@app.route('/supervisors/<int:sup_id>/continue', methods=['GET'])
@login_required
def continue_sup_draft(sup_id):
    db = get_db()
    sup = db.execute("SELECT * FROM supervisors WHERE id=?", (sup_id,)).fetchone()
    companies = db.execute("SELECT * FROM companies WHERE status='Approved' ORDER BY name").fetchall()
    db.close()
    if not sup:
        return redirect(url_for('supervisors'))
    return render_template('add_supervisor.html', draft=sup, companies=companies)

# ── SUPERVISORS ──────────────────────────────────────────
@app.route('/supervisors')
@login_required
def supervisors():
    db = get_db()
    sups = db.execute('''
        SELECT s.*, c.name as company_name
        FROM supervisors s
        LEFT JOIN companies c ON s.company_id = c.id
        ORDER BY s.created_at DESC
    ''').fetchall()
    companies = db.execute("SELECT * FROM companies WHERE status='Approved' ORDER BY name").fetchall()
    db.close()
    return render_template('supervisors.html', supervisors=sups, companies=companies)

@app.route('/supervisors/add', methods=['GET', 'POST'])
@login_required
def add_supervisor():
    if request.method == 'POST':
        if request.form.get('step_completed') != '1':
            flash('Please complete the form before submitting.', 'error')
            return redirect(url_for('add_supervisor'))
        db = get_db()
        interview = {}
        for key in request.form:
            if key.startswith('int_'):
                interview[key] = request.form[key]
        db.execute('''INSERT INTO supervisors
            (company_id, first_name, middle_name, last_name, dob, gender,
             mobile, aadhar, address, languages, qualifications, work_type,
             num_workers, experience_years, experience_months,
             technical_qualification, jobs_handled, risk_level, nature_of_work,
             interview_data, interview_result, remarks, status)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'Pending')''', (
            request.form.get('company_id'),
            request.form.get('first_name'),
            request.form.get('middle_name'),
            request.form.get('last_name'),
            request.form.get('dob'),
            request.form.get('gender'),
            request.form.get('mobile'),
            request.form.get('aadhar'),
            request.form.get('address'),
            request.form.get('languages'),
            request.form.get('qualifications'),
            request.form.get('work_type'),
            request.form.get('num_workers', 0),
            request.form.get('experience_years', 0),
            request.form.get('experience_months', 0),
            request.form.get('technical_qualification'),
            request.form.get('jobs_handled'),
            request.form.get('risk_level'),
            request.form.get('nature_of_work'),
            json.dumps(interview),
            request.form.get('interview_result'),
            request.form.get('remarks')
        ))
        sup_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]
        doc_names = request.form.getlist('doc_name[]')
        doc_files = request.files.getlist('doc_file[]')
        for name, file in zip(doc_names, doc_files):
            if file and file.filename:
                path = save_file(file, 'supervisors')
                if path:
                    db.execute('INSERT INTO supervisor_docs (supervisor_id, doc_name, file_path) VALUES (?,?,?)',
                               (sup_id, name, path))
        db.commit()
        db.close()
        flash('Supervisor submitted for approval!', 'success')
        return redirect(url_for('supervisors'))

    db = get_db()
    companies = db.execute("SELECT * FROM companies WHERE status='Approved' ORDER BY name").fetchall()
    db.close()
    return render_template('add_supervisor.html', companies=companies, draft=None)

@app.route('/supervisors/<int:sup_id>/approve', methods=['POST'])
@login_required
def approve_supervisor(sup_id):
    db = get_db()
    db.execute("UPDATE supervisors SET status='Approved' WHERE id=?", (sup_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

# ── WORKER CARDS ─────────────────────────────────────────
@app.route('/cards')
@login_required
def cards():
    db = get_db()
    company_filter = request.args.get('company', '')
    status_filter = request.args.get('status', '')
    date_filter = request.args.get('date', '')

    query = '''SELECT wc.*, c.name as company_name
               FROM worker_cards wc
               LEFT JOIN companies c ON wc.company_id = c.id
               WHERE 1=1'''
    params = []
    if company_filter:
        query += ' AND wc.company_id = ?'
        params.append(company_filter)
    if status_filter:
        query += ' AND wc.status = ?'
        params.append(status_filter)
    if date_filter:
        query += ' AND DATE(wc.created_at) = ?'
        params.append(date_filter)
    query += ' ORDER BY wc.created_at DESC'

    cards = db.execute(query, params).fetchall()
    companies = db.execute("SELECT * FROM companies WHERE status='Approved'").fetchall()
    db.close()
    return render_template('cards.html', cards=cards, companies=companies,
                           company_filter=company_filter, status_filter=status_filter,
                           date_filter=date_filter)

@app.route('/cards/add', methods=['GET', 'POST'])
@login_required
def add_card():
    if request.method == 'POST':
        db = get_db()
        company_id = request.form.get('company_id')
        card_number = generate_card_number(company_id)

        # Save photo if captured
        photo_path = None
        photo_data = request.form.get('photo_data', '')
        if photo_data and photo_data.startswith('data:image'):
            import base64
            header, data = photo_data.split(',', 1)
            photo_bytes = base64.b64decode(data)
            photo_folder = os.path.join(app.config['UPLOAD_FOLDER'], 'photos')
            os.makedirs(photo_folder, exist_ok=True)
            photo_filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{card_number.replace('/','_')}.jpg"
            photo_full_path = os.path.join(photo_folder, photo_filename)
            with open(photo_full_path, 'wb') as f:
                f.write(photo_bytes)
            photo_path = f"photos/{photo_filename}"

        # Calculate dates
        medical_valid = None
        induction_valid = None
        try:
            from dateutil.relativedelta import relativedelta
            now = datetime.now()
            medical_valid = (now + relativedelta(months=3)).strftime('%d.%b.%Y')
            induction_valid = (now + relativedelta(months=3)).strftime('%d.%b.%Y')
        except:
            from datetime import timedelta
            now = datetime.now()
            medical_valid = (now + timedelta(days=90)).strftime('%d.%b.%Y')
            induction_valid = (now + timedelta(days=90)).strftime('%d.%b.%Y')

        db.execute('''INSERT INTO worker_cards
            (card_number, company_id, supervisor_id, worker_type,
             first_name, middle_name, last_name, dob, aadhar,
             blood_group, emergency_contact, designation,
             photo_path, medical_valid, induction_valid)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
            card_number,
            company_id,
            request.form.get('supervisor_id') or None,
            request.form.get('worker_type', 'workman'),
            request.form.get('first_name'),
            request.form.get('middle_name'),
            request.form.get('last_name'),
            request.form.get('dob'),
            request.form.get('aadhar'),
            request.form.get('blood_group'),
            request.form.get('emergency_contact'),
            request.form.get('designation'),
            photo_path,
            medical_valid,
            induction_valid
        ))
        worker_id = db.execute('SELECT last_insert_rowid()').fetchone()[0]

        # Standard 4 documents
        std_docs = [
            ('Medical Certificate', 'doc_medical'),
            ('Aadhar Card', 'doc_aadhar'),
            ('Blood Report', 'doc_blood'),
            ('ESIC Card', 'doc_esic'),
        ]
        for doc_name, field in std_docs:
            file = request.files.get(field)
            if file and file.filename:
                path = save_file(file, 'workers')
                if path:
                    db.execute('INSERT INTO worker_docs (worker_id, doc_name, file_path) VALUES (?,?,?)',
                               (worker_id, doc_name, path))

        # Extra documents
        doc_names = request.form.getlist('doc_name[]')
        doc_files = request.files.getlist('doc_file[]')
        for name, file in zip(doc_names, doc_files):
            if file and file.filename:
                path = save_file(file, 'workers')
                if path:
                    db.execute('INSERT INTO worker_docs (worker_id, doc_name, file_path) VALUES (?,?,?)',
                               (worker_id, name, path))

        db.commit()
        db.close()
        flash(f'Worker card {card_number} created successfully!', 'success')
        return redirect(url_for('cards'))

    db = get_db()
    companies = db.execute("SELECT * FROM companies WHERE status='Approved' ORDER BY name").fetchall()
    supervisors = db.execute('''SELECT s.*, c.name as company_name
                                FROM supervisors s
                                LEFT JOIN companies c ON s.company_id=c.id
                                WHERE s.status='Approved' ''').fetchall()
    db.close()
    return render_template('add_card.html', companies=companies, supervisors=supervisors)

@app.route('/cards/<int:card_id>/print')
@login_required
def print_card(card_id):
    db = get_db()
    card = db.execute('''SELECT wc.*, c.name as company_name
                         FROM worker_cards wc
                         LEFT JOIN companies c ON wc.company_id=c.id
                         WHERE wc.id=?''', (card_id,)).fetchone()
    if not card:
        db.close()
        return "Card not found", 404

    # Get photo as base64 if exists
    photo_b64 = None
    if card['photo_path']:
        photo_full = os.path.join(app.config['UPLOAD_FOLDER'], card['photo_path'])
        if os.path.exists(photo_full):
            import base64
            with open(photo_full, 'rb') as f:
                photo_b64 = base64.b64encode(f.read()).decode('utf-8')

    db.close()
    return render_template('print_card.html', card=card, photo_b64=photo_b64,
                           now=datetime.now().strftime('%d.%b.%Y'))


@app.route('/cards/<int:card_id>/safety-approval', methods=['GET', 'POST'])
@login_required
def safety_approval(card_id):
    db = get_db()
    card = db.execute('''SELECT wc.*, c.name as company_name
                         FROM worker_cards wc
                         LEFT JOIN companies c ON wc.company_id=c.id
                         WHERE wc.id=?''', (card_id,)).fetchone()
    if not card:
        db.close()
        flash('Card not found', 'error')
        return redirect(url_for('cards'))

    if request.method == 'POST':
        if card['safety_status'] == 'Approved':
            flash('Card already approved and locked.', 'error')
            db.close()
            return redirect(url_for('safety_approval', card_id=card_id))
        modules = request.form.getlist('induction_modules')
        safety_name = request.form.get('safety_name', '')
        safety_remarks = request.form.get('safety_remarks', '')
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        approver = session.get('full_name') or session.get('username')
        db.execute("""UPDATE worker_cards SET
            safety_status='Approved', safety_remarks=?,
            safety_approved_by=?, safety_approved_at=?,
            induction_modules=?, status='Approved'
            WHERE id=?""", (
            safety_remarks, safety_name or approver,
            now, json.dumps(modules), card_id))
        db.commit()
        db.close()
        flash('Safety approval completed. Card is now locked.', 'success')
        return redirect(url_for('cards'))

    docs = db.execute("SELECT * FROM worker_docs WHERE worker_id=?", (card_id,)).fetchall()
    photo_b64 = None
    if card['photo_path']:
        photo_full = os.path.join(app.config['UPLOAD_FOLDER'], card['photo_path'])
        if os.path.exists(photo_full):
            import base64
            with open(photo_full, 'rb') as f:
                photo_b64 = base64.b64encode(f.read()).decode('utf-8')
    existing_modules = []
    try:
        existing_modules = json.loads(card['induction_modules'] or '[]')
    except: pass
    db.close()
    return render_template('safety_approval.html', card=card, docs=docs,
                           photo_b64=photo_b64, existing_modules=existing_modules)

@app.route('/cards/<int:card_id>/reinduction', methods=['GET', 'POST'])
@login_required
def reinduction(card_id):
    db = get_db()
    card = db.execute("SELECT * FROM worker_cards WHERE id=?", (card_id,)).fetchone()
    if not card:
        db.close()
        return redirect(url_for('cards'))
    if request.method == 'POST':
        new_valid_till = request.form.get('valid_till')
        new_induction_date = request.form.get('induction_date')
        now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        approver = session.get('full_name') or session.get('username')
        db.execute("""UPDATE worker_cards SET
            induction_valid=?, reinduction_date=?,
            reinduction_by=?, reinduction_at=?,
            safety_status='Pending', status='Pending'
            WHERE id=?""", (new_valid_till, new_induction_date, approver, now, card_id))
        db.commit()
        db.close()
        flash('Re-induction recorded. Card sent for safety approval.', 'success')
        return redirect(url_for('safety_approval', card_id=card_id))
    db.close()
    return render_template('reinduction.html', card=card,
                           today=datetime.now().strftime('%Y-%m-%d'))

# Public worker profile — no login, for QR scanning
@app.route('/worker/<card_number>')
def worker_profile(card_number):
    db = get_db()
    card = db.execute('''SELECT wc.*, c.name as company_name
                         FROM worker_cards wc
                         LEFT JOIN companies c ON wc.company_id=c.id
                         WHERE wc.card_number=?''', (card_number,)).fetchone()
    if not card:
        db.close()
        return "<h2 style='font-family:sans-serif;text-align:center;margin-top:60px;'>Worker not found</h2>", 404
    docs = db.execute("SELECT * FROM worker_docs WHERE worker_id=?", (card['id'],)).fetchall()
    photo_b64 = None
    if card['photo_path']:
        photo_full = os.path.join(app.config['UPLOAD_FOLDER'], card['photo_path'])
        if os.path.exists(photo_full):
            import base64
            with open(photo_full, 'rb') as f:
                photo_b64 = base64.b64encode(f.read()).decode('utf-8')
    modules = []
    try:
        modules = json.loads(card['induction_modules'] or '[]')
    except: pass
    db.close()
    return render_template('worker_profile.html', card=card, docs=docs,
                           photo_b64=photo_b64, modules=modules)

@login_required
def approve_card_route(card_id):
    pass

@app.route('/cards/<int:card_id>/approve', methods=['POST'])
@login_required
def approve_card(card_id):
    approval_type = request.json.get('type', 'both')
    db = get_db()
    if approval_type == 'hr':
        db.execute("UPDATE worker_cards SET hr_status='Approved', hr_remarks=? WHERE id=?",
                   (request.json.get('remarks', 'OK'), card_id))
    elif approval_type == 'safety':
        db.execute("UPDATE worker_cards SET safety_status='Approved', safety_remarks=? WHERE id=?",
                   (request.json.get('remarks', 'OK'), card_id))
    else:
        db.execute("""UPDATE worker_cards SET
                      hr_status='Approved', safety_status='Approved', status='Approved'
                      WHERE id=?""", (card_id,))
    # Check if both approved
    card = db.execute("SELECT * FROM worker_cards WHERE id=?", (card_id,)).fetchone()
    if card['hr_status'] == 'Approved' and card['safety_status'] == 'Approved':
        db.execute("UPDATE worker_cards SET status='Approved' WHERE id=?", (card_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/cards/<int:card_id>/reject', methods=['POST'])
@login_required
def reject_card(card_id):
    db = get_db()
    db.execute("UPDATE worker_cards SET status='Rejected' WHERE id=?", (card_id,))
    db.commit()
    db.close()
    return jsonify({'success': True})

@app.route('/cards/bulk-approve', methods=['POST'])
@login_required
def bulk_approve():
    ids = request.json.get('ids', [])
    db = get_db()
    for card_id in ids:
        db.execute("""UPDATE worker_cards SET
                      hr_status='Approved', safety_status='Approved', status='Approved'
                      WHERE id=?""", (card_id,))
    db.commit()
    db.close()
    return jsonify({'success': True, 'count': len(ids)})

# ── API: Get supervisors by company ──────────────────────
@app.route('/api/supervisors/<int:company_id>')
@login_required
def api_supervisors(company_id):
    db = get_db()
    sups = db.execute("SELECT id, first_name, last_name FROM supervisors WHERE company_id=? AND status='Approved'",
                      (company_id,)).fetchall()
    db.close()
    return jsonify([{'id': s['id'], 'name': f"{s['first_name']} {s['last_name']}"} for s in sups])

# ── REPORTS ──────────────────────────────────────────────
@app.route('/reports')
@login_required
def reports():
    db = get_db()
    companies = db.execute("SELECT * FROM companies").fetchall()
    db.close()
    return render_template('reports.html', companies=companies)

@app.route('/reports/data')
@login_required
def reports_data():
    db = get_db()
    company_id = request.args.get('company', '')
    work_type = request.args.get('work_type', '')
    status = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')

    query = '''SELECT wc.*, c.name as company_name
               FROM worker_cards wc
               LEFT JOIN companies c ON wc.company_id = c.id
               WHERE 1=1'''
    params = []
    if company_id:
        query += ' AND wc.company_id=?'; params.append(company_id)
    if status:
        query += ' AND wc.status=?'; params.append(status)
    if date_from:
        query += ' AND DATE(wc.created_at)>=?'; params.append(date_from)
    if date_to:
        query += ' AND DATE(wc.created_at)<=?'; params.append(date_to)
    query += ' ORDER BY wc.created_at DESC'

    cards = db.execute(query, params).fetchall()
    db.close()
    return jsonify([dict(c) for c in cards])

@app.route('/reports/export/excel')
@login_required
def export_excel():
    db = get_db()
    company_id = request.args.get('company', '')
    status = request.args.get('status', '')
    date_from = request.args.get('date_from', '')
    date_to = request.args.get('date_to', '')

    query = '''SELECT wc.card_number, wc.first_name, wc.middle_name, wc.last_name,
                      wc.dob, wc.blood_group, wc.designation, wc.emergency_contact,
                      c.name as company_name, wc.status, wc.created_at
               FROM worker_cards wc
               LEFT JOIN companies c ON wc.company_id = c.id
               WHERE 1=1'''
    params = []
    if company_id:
        query += ' AND wc.company_id=?'; params.append(company_id)
    if status:
        query += ' AND wc.status=?'; params.append(status)
    if date_from:
        query += ' AND DATE(wc.created_at)>=?'; params.append(date_from)
    if date_to:
        query += ' AND DATE(wc.created_at)<=?'; params.append(date_to)

    cards = db.execute(query, params).fetchall()
    db.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Worker Cards'

    # Header style
    header_fill = PatternFill("solid", fgColor="C8102E")
    header_font = Font(bold=True, color="FFFFFF")

    headers = ['Card No.', 'First Name', 'Middle Name', 'Last Name', 'DOB',
               'Blood Group', 'Designation', 'Emergency Contact',
               'Company', 'Status', 'Date Created']

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row, card in enumerate(cards, 2):
        vals = [card['card_number'],
                card['first_name'], card['middle_name'], card['last_name'],
                card['dob'], card['blood_group'], card['designation'],
                card['emergency_contact'], card['company_name'],
                card['status'], card['created_at']]
        for col, val in enumerate(vals, 1):
            ws.cell(row=row, column=col, value=val)

    # Auto width
    for col in ws.columns:
        max_len = max((len(str(cell.value or '')) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"CMC_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output, download_name=filename,
                     as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/uploads/<path:filename>')
@login_required
def serve_upload(filename):
    return send_file(os.path.join(app.config['UPLOAD_FOLDER'], filename))

# ── SETUP: Create first admin user ───────────────────────
@app.route('/setup', methods=['GET', 'POST'])
def setup():
    db = get_db()
    existing = db.execute("SELECT COUNT(*) FROM users").fetchone()[0]
    if existing > 0:
        db.close()
        return redirect(url_for('login'))
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        full_name = request.form.get('full_name')
        db.execute('INSERT INTO users (username, password_hash, full_name, role) VALUES (?,?,?,?)',
                   (username, generate_password_hash(password), full_name, 'admin'))
        db.commit()
        db.close()
        return redirect(url_for('login'))
    db.close()
    return render_template('setup.html')

if __name__ == '__main__':
    init_db()
    from database import migrate_db
    migrate_db()
    # Get local IP for network access
    import socket
    hostname = socket.gethostname()
    local_ip = socket.gethostbyname(hostname)
    print(f"\n{'='*50}")
    print(f"  CMS Server Starting...")
    print(f"  Local:   http://localhost:5000")
    print(f"  Network: http://{local_ip}:5000")
    print(f"  Share the Network URL with other computers")
    print(f"{'='*50}\n")
    app.run(host='0.0.0.0', port=5000, debug=False)
